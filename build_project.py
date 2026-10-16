"""
Sales Analytics Project - Excel + Power BI
Run:
    pip install pandas openpyxl numpy
    python build_project.py
This script reads sales_data.csv and regenerates the Excel workbook.
"""
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
CSV_FILE = ROOT / "sales_data.csv"
OUTPUT_FILE = ROOT / "Sales_Analytics_Dashboard.xlsx"

def main():
    df = pd.read_csv(CSV_FILE, parse_dates=["Order_Date"])
    required = {"Order_ID", "Order_Date", "Region", "Category", "Product",
                "Units_Sold", "Sales_INR", "Profit_INR"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Missing required columns: {sorted(missing)}")

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        df.to_excel(writer, sheet_name="Sales_Data", index=False)
        df.groupby("Region", as_index=False).agg(
            Total_Sales=("Sales_INR", "sum"),
            Total_Profit=("Profit_INR", "sum"),
            Orders=("Order_ID", "count")
        ).to_excel(writer, sheet_name="Region_Summary", index=False)
        df.groupby("Category", as_index=False).agg(
            Total_Sales=("Sales_INR", "sum"),
            Total_Profit=("Profit_INR", "sum"),
            Units_Sold=("Units_Sold", "sum")
        ).to_excel(writer, sheet_name="Category_Summary", index=False)
        df.assign(Month=df["Order_Date"].dt.to_period("M").astype(str)).groupby(
            "Month", as_index=False
        ).agg(
            Total_Sales=("Sales_INR", "sum"),
            Total_Profit=("Profit_INR", "sum")
        ).to_excel(writer, sheet_name="Monthly_Trend", index=False)

    wb = load_workbook(OUTPUT_FILE)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col[:200])
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 2, 12), 24)

    ws = wb["Region_Summary"]
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Total Sales by Region"
    chart.y_axis.title = "Region"
    chart.x_axis.title = "Sales (INR)"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(chart, "E2")

    ws = wb["Monthly_Trend"]
    chart = LineChart()
    chart.style = 13
    chart.title = "Monthly Sales Trend"
    chart.y_axis.title = "Sales (INR)"
    chart.x_axis.title = "Month"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    ws.add_chart(chart, "E2")
    wb.save(OUTPUT_FILE)
    print(f"Created: {OUTPUT_FILE}")

if __name__ == "__main__":
    main()
